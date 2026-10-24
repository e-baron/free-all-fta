import importlib.util
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

import openpyxl

SCRIPT = Path(__file__).resolve().parents[1] / "free-all-fta.py"
spec = importlib.util.spec_from_file_location("free_all_fta", SCRIPT)
fta = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = fta
spec.loader.exec_module(fta)


def workbook_for(specification, models):
    wb = openpyxl.Workbook()
    events = wb.active
    events.title = "events"
    events.append(["id", "model_type", "allocated value", "mean_repair_time"])
    for event_id, model_type in models.items():
        events.append([event_id, model_type, 0.001, 2 if model_type in ("ConstantRate", "Fixed") else None])
    gates = wb.create_sheet("gates")
    gates.append(["id", "formula", "safety target", "calculated frequency", "calculated mean repair time", "differences"])
    for gate_id in specification["gates"]:
        gates.append([gate_id, "", 1, None, None, None])
    tree = wb.create_sheet("tree")
    max_depth = max(3, specification.get("max_depth", 4))
    headers = []
    for level in range(1, max_depth + 1):
        headers.extend([f"L{level}_ID", f"L{level}_TYPE"])
    tree.append(headers)

    def add_node(node, level):
        if isinstance(node, str):
            row = [None] * len(headers)
            row[(level - 1) * 2] = node
            tree.append(row)
            return
        gate_id, gate_type, children = node
        row = [None] * len(headers)
        row[(level - 1) * 2] = gate_id
        row[(level - 1) * 2 + 1] = gate_type
        tree.append(row)
        for child in children:
            add_node(child, level + 1)

    add_node(specification["root"], 1)
    return wb


def tree(root, gates):
    return {"root": root, "gates": gates, "max_depth": 6}


class ModelValidationTests(unittest.TestCase):
    def assert_valid(self, root, gates, models):
        wb = workbook_for(tree(root, gates), models)
        try:
            fta.validate_model(wb)
        finally:
            wb.close()

    def assert_invalid(self, root, gates, models, contains=None):
        wb = workbook_for(tree(root, gates), models)
        try:
            with self.assertRaises(ValueError) as error:
                fta.validate_model(wb)
            if contains:
                self.assertIn(contains, str(error.exception))
        finally:
            wb.close()

    def assert_classification(self, root, gates, models, expected):
        wb = workbook_for(tree(root, gates), models)
        try:
            validated = fta.validate_model(wb)
            gate_id = root[0]
            self.assertEqual(validated["classifications"][gate_id].classification, expected)
        finally:
            wb.close()

    def test_semantic_classifications_propagate_through_or_and(self):
        C = fta.Classification
        cases = [
            (("G", "OR", ["R1", "R2"]), {"G"}, {"R1": "ConstantRate", "R2": "ConstantRate"}, C.FREQUENCY),
            (("G", "OR", ["F1", "F2"]), {"G"}, {"F1": "Fixed", "F2": "Fixed"}, C.FREQUENCY),
            (("G", "OR", ["F", "R"]), {"G"}, {"F": "Fixed", "R": "ConstantRate"}, C.FREQUENCY),
            (("G", "OR", [("N", "OR", ["F", "R"]), "R2"]), {"G", "N"}, {"F": "Fixed", "R": "ConstantRate", "R2": "ConstantRate"}, C.FREQUENCY),
            (("G", "OR", ["P1", "P2"]), {"G"}, {"P1": "Probability", "P2": "Probability"}, C.DIMENSIONLESS),
            (("G", "AND", ["R1", "R2"]), {"G"}, {"R1": "ConstantRate", "R2": "ConstantRate"}, C.FREQUENCY),
            (("G", "AND", ["F", "P"]), {"G"}, {"F": "Fixed", "P": "Probability"}, C.FREQUENCY),
            (("G", "AND", [("N", "OR", ["F", "R"]), "P"]), {"G", "N"}, {"F": "Fixed", "R": "ConstantRate", "P": "Probability"}, C.FREQUENCY),
            (("G", "AND", [("N", "OR", ["F1", "F2"]), "P"]), {"G", "N"}, {"F1": "Fixed", "F2": "Fixed", "P": "Probability"}, C.FREQUENCY),
            (("G", "AND", ["P1", "P2"]), {"G"}, {"P1": "Probability", "P2": "Probability"}, C.DIMENSIONLESS),
        ]
        for root, gates, models, expected in cases:
            with self.subTest(root=root, expected=expected):
                self.assert_classification(root, gates, models, expected)

    def test_nested_frequency_gate_can_join_another_frequency_with_mttr(self):
        root = ("TOP", "AND", [("MIX", "OR", ["F", "R1"]), "R2"])
        self.assert_classification(root, {"TOP", "MIX"}, {"F": "Fixed", "R1": "ConstantRate", "R2": "ConstantRate"}, fta.Classification.FREQUENCY)

    def test_numeric_values_are_validated_when_available(self):
        invalid_cases = [
            ("Probability", 1.2, None, "between 0 and 1"),
            ("Multiplicity", -1, None, "non-negative"),
            ("ConstantRate", -0.1, 2, "non-negative"),
            ("ConstantRate", 0.1, 0, "strictly positive"),
        ]
        for event_model, value, duration, expected in invalid_cases:
            wb = workbook_for(tree(("G", "AND", ["E"]), {"G"}), {"E": event_model})
            wb["events"]["C2"] = value
            wb["events"]["D2"] = duration
            try:
                with self.assertRaisesRegex(ValueError, expected):
                    fta.validate_model(wb)
            finally:
                wb.close()

    def test_formula_based_inputs_are_evaluated_or_rejected(self):
        for formula, should_pass in (("=0.002", True), ("=1/0", False)):
            wb = workbook_for(tree(("G", "AND", ["F"]), {"G"}), {"F": "Fixed"})
            wb["events"]["C2"] = formula
            wb["events"]["D2"] = "=2"
            # Use a saved workbook so formula and cached-value loading follows the
            # same path as CLI validation.
            with tempfile.TemporaryDirectory() as tmp:
                source = Path(tmp) / "formula.xlsx"
                wb.save(source)
                wb.close()
                formula_wb = openpyxl.load_workbook(source, data_only=False)
                values_wb = openpyxl.load_workbook(source, data_only=True)
                try:
                    if should_pass:
                        validated = fta.validate_model(formula_wb, value_workbook=values_wb)
                        self.assertTrue(validated["classifications"]["F"].equivalent_restoration_time_available)
                    else:
                        with self.assertRaisesRegex(ValueError, "could not be evaluated"):
                            fta.validate_model(formula_wb, value_workbook=values_wb)
                finally:
                    formula_wb.close()
                    values_wb.close()

    def test_permitted_gate_combinations(self):
        cases = [
            (("G", "AND", ["R1", "R2"]), {"G"}, {"R1": "ConstantRate", "R2": "ConstantRate"}),
            (("G", "AND", ["R1", "P"]), {"G"}, {"R1": "ConstantRate", "P": "Probability"}),
            (("G", "AND", ["R1", "R2", "P"]), {"G"}, {"R1": "ConstantRate", "R2": "ConstantRate", "P": "Probability"}),
            (("G", "AND", ["F", "P"]), {"G"}, {"F": "Fixed", "P": "Probability"}),
            (("G", "AND", ["F", "P", "M"]), {"G"}, {"F": "Fixed", "P": "Probability", "M": "Multiplicity"}),
            (("G", "AND", ["F1", "F2"]), {"G"}, {"F1": "Fixed", "F2": "Fixed"}),
            (("G", "AND", ["F", "R"]), {"G"}, {"F": "Fixed", "R": "ConstantRate"}),
            (("G", "AND", ["F", "R", "P"]), {"G"}, {"F": "Fixed", "R": "ConstantRate", "P": "Probability"}),
            (("G", "AND", ["P1", "P2"]), {"G"}, {"P1": "Probability", "P2": "Probability"}),
            (("G", "AND", ["P", "M"]), {"G"}, {"P": "Probability", "M": "Multiplicity"}),
            (("G", "AND", ["M", "T"]), {"G"}, {"M": "Multiplicity", "T": "True"}),
            (("G", "AND", ["T", "N"]), {"G"}, {"T": "True", "N": "False"}),
            (("G", "OR", ["F", "R"]), {"G"}, {"F": "Fixed", "R": "ConstantRate"}),
            (("G", "OR", ["F1", "F2"]), {"G"}, {"F1": "Fixed", "F2": "Fixed"}),
            (("G", "OR", ["R1", "R2"]), {"G"}, {"R1": "ConstantRate", "R2": "ConstantRate"}),
            (("G", "OR", [("N", "AND", ["R1", "R2"]), "F"]), {"G", "N"}, {"R1": "ConstantRate", "R2": "ConstantRate", "F": "Fixed"}),
        ]
        for root, gates, models in cases:
            with self.subTest(root=root):
                self.assert_valid(root, gates, models)

    def test_forbidden_gate_combinations(self):
        cases = [
            (("G", "OR", ["P", "R"]), {"G"}, {"P": "Probability", "R": "ConstantRate"}, "cannot be combined by OR"),
            (("G", "OR", ["P", "F"]), {"G"}, {"P": "Probability", "F": "Fixed"}, "cannot be combined by OR"),
            (("G", "OR", ["M", "P"]), {"G"}, {"M": "Multiplicity", "P": "Probability"}, "AND-only"),
            (("G", "OR", ["M", "R"]), {"G"}, {"M": "Multiplicity", "R": "ConstantRate"}, "AND-only"),
            (("G", "AND", ["U"]), {"G"}, {"U": "Mystery"}, "unknown model"),
            (("G", "OR", ["U", "F"]), {"G"}, {"U": "Mystery", "F": "Fixed"}, "unknown model"),
            (("G", None, ["R"]), {"G"}, {"R": "ConstantRate"}, "no recognized AND/OR type"),
        ]
        for root, gates, models, expected in cases:
            with self.subTest(root=root, expected=expected):
                self.assert_invalid(root, gates, models, expected)

    def test_multiple_frequency_and_requires_valid_restoration_times(self):
        wb = workbook_for(tree(("G", "AND", ["R1", "R2"]), {"G"}), {"R1": "ConstantRate", "R2": "ConstantRate"})
        wb["events"]["D2"] = None
        try:
            with self.assertRaisesRegex(ValueError, "MTTR"):
                fta.validate_model(wb)
        finally:
            wb.close()

    def test_fixed_requires_valid_restoration_time(self):
        for mttr in (None, 0, -1):
            wb = workbook_for(tree(("G", "OR", ["F", "R"]), {"G"}), {"F": "Fixed", "R": "ConstantRate"})
            wb["events"]["D2"] = mttr
            try:
                with self.assertRaisesRegex(ValueError, "Fixed event.*MTTR|MTTR"):
                    fta.validate_model(wb)
            finally:
                wb.close()

    def test_cut_set_expansion_is_validated(self):
        root = ("TOP", "AND", [("MIX", "OR", ["F", "R1"]), "R2"])
        wb = workbook_for(tree(root, {"TOP", "MIX"}), {"F": "Fixed", "R1": "ConstantRate", "R2": "ConstantRate"})
        try:
            fta.validate_model(wb)
        finally:
            wb.close()
        root = ("TOP", "AND", [("MIX", "OR", ["F1", "F2", "R1"]), "R2"])
        self.assert_valid(root, {"TOP", "MIX"}, {"F1": "Fixed", "F2": "Fixed", "R1": "ConstantRate", "R2": "ConstantRate"})

    def test_cut_set_validator_rejects_unknown_event_models(self):
        event_classes = {
            "U": fta.ClassificationResult("U", fta.Classification.INVALID, "Mystery"),
        }
        errors = fta.validate_cut_sets([("U",)], event_classes.__getitem__, {}, {})
        self.assertEqual(len(errors), 1)
        self.assertIn("Minimal cut set CS-1", errors[0])
        self.assertIn("Mystery", errors[0])

    def test_cut_set_validator_reuses_and_semantics(self):
        C = fta.Classification
        events = {
              "R1": fta.ClassificationResult("R1", C.FREQUENCY, "ConstantRate", has_frequency=True, equivalent_restoration_time_available=True, frequency_value=0.1, mttr_value=2),
              "R2": fta.ClassificationResult("R2", C.FREQUENCY, "ConstantRate", has_frequency=True, equivalent_restoration_time_available=True, frequency_value=0.2, mttr_value=3),
              "F1": fta.ClassificationResult("F1", C.FREQUENCY, "Fixed", has_frequency=True, equivalent_restoration_time_available=True, frequency_value=0.1, mttr_value=1),
              "F2": fta.ClassificationResult("F2", C.FREQUENCY, "Fixed", has_frequency=True, equivalent_restoration_time_available=True, frequency_value=0.2, mttr_value=1),
              "P": fta.ClassificationResult("P", C.DIMENSIONLESS, "Probability", factor_value=0.2),
        }
        classify = events.__getitem__
        self.assertEqual(fta.validate_cut_sets([("R1", "R2"), ("F1", "P")], classify, {}, {}), [])
        self.assertEqual(fta.validate_cut_sets(
            [("F1", "F2"), ("F1", "R1"), ("F1", "R1", "P")], classify, {}, {}
        ), [])
        missing_mttr = fta.ClassificationResult("F3", C.FREQUENCY, "Fixed", has_frequency=True)
        errors = fta.validate_cut_sets([("F1", "F3")], lambda eid: missing_mttr if eid == "F3" else classify(eid), {}, {})
        self.assertEqual(len(errors), 1)
        self.assertIn("CS-1", errors[0])

    def test_frequency_formulas_and_weighted_mttr_match_markov_equations(self):
        C = fta.Classification
        rate_a = fta.ClassificationResult("A", C.FREQUENCY, "ConstantRate", has_frequency=True,
                                          equivalent_restoration_time_available=True, frequency_value=0.1, mttr_value=2)
        rate_b = fta.ClassificationResult("B", C.FREQUENCY, "Fixed", has_frequency=True,
                                          equivalent_restoration_time_available=True, frequency_value=0.2, mttr_value=3)
        probability = fta.ClassificationResult("Q", C.DIMENSIONLESS, "Probability", factor_value=0.25)
        combined = fta._validate_and("AND", [rate_a, rate_b, probability])
        self.assertAlmostEqual(combined.frequency_value, 0.025)
        self.assertAlmostEqual(combined.mttr_value, 1.2)
        combined_fixed = fta._validate_and("AND", [rate_a, rate_b])
        self.assertAlmostEqual(combined_fixed.frequency_value, 0.1)
        self.assertAlmostEqual(combined_fixed.mttr_value, 1.2)
        summed = fta._validate_or("OR", [rate_a, rate_b])
        self.assertAlmostEqual(summed.frequency_value, 0.3)
        self.assertAlmostEqual(summed.mttr_value, (0.1 * 2 + 0.2 * 3) / 0.3)

    def test_zero_frequency_or_has_no_equivalent_mttr(self):
        C = fta.Classification
        zero_a = fta.ClassificationResult("A", C.FREQUENCY, "Fixed", has_frequency=True,
                                          equivalent_restoration_time_available=True, frequency_value=0, mttr_value=2)
        zero_b = fta.ClassificationResult("B", C.FREQUENCY, "ConstantRate", has_frequency=True,
                                          equivalent_restoration_time_available=True, frequency_value=0, mttr_value=3)
        result = fta._validate_or("OR", [zero_a, zero_b])
        self.assertEqual(result.frequency_value, 0)
        self.assertIsNone(result.mttr_value)
        self.assertFalse(result.equivalent_restoration_time_available)

    def test_invalid_cut_set_does_not_create_workbook_or_directory(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "valid_tree.xlsx"
            wb = workbook_for(tree(("G", "OR", ["F", "R"]), {"G"}), {"F": "Fixed", "R": "ConstantRate"})
            wb.save(source)
            wb.close()
            out_dir = root / "must-not-exist"
            with patch.object(fta, "minimal_cut_sets", return_value=[("F", "MISSING")]):
                with self.assertRaisesRegex(ValueError, "Minimal cut set CS-1"):
                    fta.build_markov_workbook(str(source), str(out_dir / "result.xlsx"))
            self.assertFalse(out_dir.exists())

    def test_fixed_or_constant_rate_generates_frequency_with_weighted_mttr(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "mixed.xlsx"
            wb = workbook_for(tree(("G", "OR", ["F", "R"]), {"G"}), {"F": "Fixed", "R": "ConstantRate"})
            wb.save(source)
            wb.close()
            output = fta.build_markov_workbook(str(source))
            result_wb = openpyxl.load_workbook(output, data_only=False)
            try:
                gate = result_wb["gates"]
                columns = fta.headers(gate)
                self.assertTrue(gate.cell(2, columns["calculated frequency"]).value.startswith("="))
                self.assertTrue(gate.cell(2, columns["calculated mean repair time"]).value.startswith("=IFERROR"))
            finally:
                result_wb.close()

    def test_true_or_operand_is_boolean_true_and_minimizes_to_empty_set(self):
        wb = workbook_for(tree(("G", "OR", ["T", "P"]), {"G"}), {"T": "True", "P": "Probability"})
        try:
            result = fta.validate_model(wb)
            self.assertEqual(result["classifications"]["G"].classification, fta.Classification.BOOLEAN_TRUE)
            self.assertEqual(result["cut_sets"], [()])
        finally:
            wb.close()

    def test_nested_boolean_gate_uses_dimensionless_calculation(self):
        root = ("TOP", "AND", [("BOOL", "AND", ["T", "F"]), "R"])
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "boolean.xlsx"
            wb = workbook_for(tree(root, {"TOP", "BOOL"}), {"T": "True", "F": "False", "R": "ConstantRate"})
            wb.save(source)
            wb.close()
            output = fta.build_markov_workbook(str(source))
            result_wb = openpyxl.load_workbook(output, data_only=False)
            try:
                gates = result_wb["gates"]
                columns = fta.headers(gates)
                self.assertEqual(gates.cell(2, columns["calculated frequency"]).value, "=0")
                self.assertEqual(gates.cell(3, columns["calculated probability"]).value, "=0")
            finally:
                result_wb.close()

    def test_invalid_cli_inputs_do_not_create_or_overwrite_outputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root_dir = Path(tmp)
            source = root_dir / "invalid.xlsx"
            wb = workbook_for(tree(("G", "OR", ["F", "P"]), {"G"}), {"F": "Fixed", "P": "Probability"})
            wb.save(source)
            wb.close()
            output_dir = root_dir / "not-created"
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--diag", "--diag-dir", str(output_dir)],
                capture_output=True, text=True,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(output_dir.exists())
            self.assertEqual(list(root_dir.glob("*_alloc.xlsx")), [])
            self.assertEqual(list(root_dir.glob("*.svg")), [])
            self.assertEqual(list(root_dir.glob("*.html")), [])

            existing_dir = root_dir / "existing"
            existing_dir.mkdir()
            output = existing_dir / "invalid_alloc.xlsx"
            output.write_bytes(b"keep existing output")
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--diag-dir", str(existing_dir)],
                capture_output=True, text=True,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(output.read_bytes(), b"keep existing output")

    def test_valid_cli_generates_workbook_and_diagrams(self):
        with tempfile.TemporaryDirectory() as tmp:
            root_dir = Path(tmp)
            source = root_dir / "valid.xlsx"
            wb = workbook_for(tree(("G", "AND", ["R1", "R2"]), {"G"}), {"R1": "ConstantRate", "R2": "ConstantRate"})
            wb.save(source)
            wb.close()
            out_dir = root_dir / "outputs"
            result = subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--diag", "--diag-dir", str(out_dir)],
                capture_output=True, text=True,
            )
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertTrue((out_dir / "valid_alloc.xlsx").exists())
            self.assertTrue((out_dir / "valid_alloc_fta.html").exists())
            self.assertTrue(list(out_dir.glob("*.svg")))

    def test_diagram_fallback_uses_validated_mttr_for_nested_formula_gates(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "nested.xlsx"
            inner = ("INNER", "OR", ["F", "R1"])
            root_node = ("TOP", "AND", [inner, "R2", "P"])
            wb = workbook_for(tree(root_node, {"TOP", "INNER"}), {
                "F": "Fixed", "R1": "ConstantRate", "R2": "ConstantRate", "P": "Probability"
            })
            wb.save(source)
            wb.close()
            output = fta.build_markov_workbook(str(source))
            diagram_dir = root / "diagrams"
            html = fta.build_diagrams(output, output_dir=diagram_dir)
            self.assertTrue(Path(html).exists())
            self.assertTrue(list(diagram_dir.glob("*.svg")))

    def test_diagram_ignores_stale_cached_gate_frequency(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "stale-cache.xlsx"
            wb = workbook_for(tree(("G", "AND", ["R", "P"]), {"G"}), {"R": "ConstantRate", "P": "Probability"})
            events = wb["events"]
            events["C3"] = 0.25
            gates = wb["gates"]
            gates["D2"] = 999
            gates["E2"] = 888
            wb.save(source)
            wb.close()

            data = fta._diagram_data(str(source))
            gate_data = data[5]["G"]
            self.assertAlmostEqual(gate_data["w"], 0.00025)
            self.assertEqual(gate_data["sdt"], 2)


if __name__ == "__main__":
    unittest.main()
